import os
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH

doc = docx.Document()

# Standard Academic Thesis Margins: 1 inch (2.54 cm) on all sides
for section in doc.sections:
    section.top_margin = Inches(1.0)
    section.bottom_margin = Inches(1.0)
    section.left_margin = Inches(1.0)
    section.right_margin = Inches(1.0)
    section.header_distance = Inches(0.5)
    section.footer_distance = Inches(0.5)

# Normal Paragraph Styling
style_normal = doc.styles['Normal']
style_normal.font.name = 'Times New Roman'
style_normal.font.size = Pt(12)
style_normal.font.color.rgb = RGBColor(30, 30, 30)
style_normal.paragraph_format.line_spacing = 1.45
style_normal.paragraph_format.space_after = Pt(8)
style_normal.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY

# Title
title_p = doc.add_paragraph()
title_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
title_p.paragraph_format.space_before = Pt(14)
title_p.paragraph_format.space_after = Pt(24)
title_run = title_p.add_run("ACKNOWLEDGEMENTS")
title_run.font.name = 'Times New Roman'
title_run.font.size = Pt(16)
title_run.font.bold = True
title_run.font.color.rgb = RGBColor(17, 24, 39)

paragraphs = [
    (
        "Submitting this doctoral dissertation marks the completion of an intensive, multi-year research endeavor dedicated "
        "to understanding and advancing materials for electrochemical energy storage. While a doctoral thesis is submitted under "
        "the name of an individual scholar, the scientific progress it documents is intrinsically collective. It is built upon the "
        "patient mentorship of senior researchers, the active cooperation of laboratory peers, the provision of institutional "
        "infrastructure, and the steadfast encouragement of family. It is with deep respect and gratitude that I record my "
        "acknowledgement to all who contributed to the realization of this work."
    ),
    (
        "First and foremost, I wish to express my profound intellectual debt and gratitude to my research supervisor and Head "
        "of the Department, Dr. C. K. Sumesh. His guidance has been the guiding force throughout my doctoral studies. From the "
        "initial conceptualization of research objectives through synthesis optimizations, material characterization, and the "
        "rigorous interpretation of electrochemical data, Dr. Sumesh provided continuous insight and constructive critique. In "
        "experimental energy storage research, where unexpected cell polarization, rapid capacity fading, or complex impedance "
        "responses frequently challenge initial assumptions, his measured approach to problem-solving taught me to analyze "
        "discrepancies systematically rather than dismiss them. He demanded scientific precision, reproducibility, and high "
        "standards of analytical reporting, while granting me the intellectual independence necessary to develop confidence as "
        "a researcher. I am deeply thankful for his accessible mentorship, his patience across every revision, and his consistent "
        "encouragement to pursue meaningful scientific questions."
    ),
    (
        "I express my sincere appreciation to Dr. Gayatri Dave, Dean, for her academic leadership, administrative support, and "
        "continuing encouragement of doctoral research scholars. Her efforts in fostering an environment conducive to interdisciplinary "
        "inquiry ensured that institutional reviews, academic evaluations, and departmental processes proceeded with clarity and "
        "efficiency throughout my tenure. Her broad perspective on scientific education and research development has been a source "
        "of steady inspiration."
    ),
    (
        "I am grateful to Charotar University of Science and Technology (CHARUSAT), Changa, Gujarat, for providing the institutional "
        "foundation, research ecosystem, and physical facilities that made this study possible. The university's central instrumentation "
        "laboratories, computational infrastructure, and extensive access to scientific journals and databases provided the tools "
        "required to carry out advanced material synthesis, structural diagnostics, and electrochemical measurements. The serene and "
        "scholarly atmosphere of the Changa campus provided an ideal setting for focused doctoral investigation."
    ),
    (
        "I acknowledge the faculty colleagues and technical staff of the Department for their professional cooperation and assistance. "
        "My special thanks go to the laboratory technical staff who helped maintain characterization equipment, ensured consistent "
        "utility supplies, and assisted with laboratory safety protocols during synthesis and testing runs."
    ),
    (
        "Experimental work in energy storage involves substantial material overheads, including high-purity chemicals, battery-grade "
        "active components, specialized separators, coin-cell hardware, and access to precision analytical testing. I gratefully "
        "acknowledge the research fellowship and funding support provided by [XYZ Funding Agency]. This financial sponsorship ensured "
        "the steady continuity of experimental cycles, covered essential characterization services, and supported the presentation of "
        "our findings at academic conferences, allowing our research to be discussed and evaluated by the wider scientific community."
    ),
    (
        "The daily reality of experimental research is shaped profoundly by the colleagues with whom one shares the laboratory. "
        "I have been fortunate to work alongside a dedicated and supportive group of researchers: Rahul, Parth, Divyam, Nandini, "
        "Pooja, Shobhraj, Samruddhi, Simmy, Vibhuti, Pratiksha, Upmanyu, Hardi, Harsh, Kinjal, and Krishna. Our shared hours in the "
        "laboratory—preparing slurries, coating electrodes, assembling coin cells under inert atmospheres, analyzing cyclic "
        "voltammetry curves, evaluating galvanostatic charge-discharge profiles, and tracking long-term cycling stability—formed the "
        "practical core of this work. Beyond the technical collaborations and equipment sharing, their camaraderie, ready humor, "
        "and willingness to assist during demanding instrument shifts made the laboratory a productive and welcoming environment. "
        "I am sincerely thankful for their scientific cooperation, thoughtful feedback, and lasting friendship."
    ),
    (
        "Behind any sustained individual achievement lies the silent, selfless foundation provided by family. I express my "
        "heartfelt reverence and gratitude to my parents, Mukesh Shah and Ranjana Shah. Their lifetime of hard work, continuous "
        "sacrifices, and quiet confidence in my abilities have sustained me through every stage of my education. They taught me the "
        "virtues of discipline, patience, and perseverance through their own lives long before I stepped into a research laboratory. "
        "Their understanding during demanding university schedules, late evening working hours, and long stretches of concentrated "
        "writing provided the emotional anchor that enabled me to persist. Every milestone recorded in this dissertation is a direct "
        "outcome of their values, prayers, and unconditional care."
    ),
    (
        "I extend my deep appreciation to my brother, Harsh Jagetiya, whose practical guidance, optimism, and companionship consistently "
        "helped me maintain perspective and renewed energy during challenging junctures of this doctoral journey. I also record my "
        "respectful gratitude to my grandparents and extended family members, whose blessings, affection, and quiet pride have always "
        "given me strength and grounded confidence."
    ),
    (
        "My deepest and most personal gratitude belongs to my spouse, Shubham Devpura (Chartered Accountant). Navigating a doctoral "
        "program is a demanding undertaking that tests the endurance of family life, and Shubham met every challenge of this journey "
        "with extraordinary patience, maturity, and emotional generosity. His analytical mindset, professional composure, and steady "
        "belief in my aspirations provided a tranquil counterbalance to the pressures of experimental timelines and thesis submission "
        "deadlines. Whether accommodating prolonged laboratory commitments, encouraging me through periods of experimental fatigue, "
        "or sharing quiet pride in every small breakthrough, his companionship has been my greatest source of calm and strength. This "
        "thesis stands as much as a testament to his understanding and partnership as it does to my scientific labor."
    ),
    (
        "Finally, I extend my appreciation to all teachers, colleagues, and well-wishers whose guidance, insightful inquiries, or "
        "encouraging words contributed in direct or indirect ways to the completion of this dissertation."
    )
]

for p_text in paragraphs:
    p = doc.add_paragraph(p_text)
    p.paragraph_format.line_spacing = 1.45
    p.paragraph_format.space_after = Pt(8)
    p.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY

doc.add_paragraph()  # spacing

# Signature Table
sig_table = doc.add_table(rows=1, cols=2)
sig_table.autofit = False
sig_table.columns[0].width = Inches(3.4)
sig_table.columns[1].width = Inches(3.1)

cell_left = sig_table.cell(0, 0)
p_left = cell_left.paragraphs[0]
p_left.paragraph_format.line_spacing = 1.2
p_left.paragraph_format.space_after = Pt(2)
r_date = p_left.add_run("Date: ______________\nPlace: Changa, Gujarat")
r_date.font.name = 'Times New Roman'
r_date.font.size = Pt(11)
r_date.font.color.rgb = RGBColor(70, 70, 70)

cell_right = sig_table.cell(0, 1)
p_right = cell_right.paragraphs[0]
p_right.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.RIGHT
p_right.paragraph_format.line_spacing = 1.2
p_right.paragraph_format.space_after = Pt(2)
r_sig = p_right.add_run("[Candidate Name / Signature]\nPh.D. Research Scholar\nCHARUSAT, Changa, Gujarat")
r_sig.font.name = 'Times New Roman'
r_sig.font.size = Pt(11)
r_sig.font.bold = True
r_sig.font.color.rgb = RGBColor(20, 20, 20)

target_path = r"C:\Users\SantoshMadnani\Documents\CA_Trader\PhD_Thesis_Acknowledgement.docx"
doc.save(target_path)
print(f"Generated Word file: {target_path}")
print(f"File size: {os.path.getsize(target_path):,} bytes")

